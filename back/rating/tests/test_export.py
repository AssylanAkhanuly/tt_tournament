"""Выгрузка рейтинг-листа в Excel ✳ (01.10.2026).

Файл открывается и читается по ячейкам: не «ответ 200 и что-то пришло», а
заголовки, значения, порядок строк и то, что отбор и сортировка экрана дошли до
файла. Выгрузка, которая отдаёт не тот список, выглядит исправной, пока файл
не откроют.
"""
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from rating import engine, services
from rating.export import fio
from rating.models import RatingParams
from users.models import User

pytestmark = pytest.mark.django_db

URL = "/api/rating/export/"


@pytest.fixture
def api():
    return APIClient()


@pytest.fixture
def params():
    return RatingParams.objects.create(d=15, k_standard=Decimal("0.6"), is_active=True)


def игрок(phone, name, legacy, **over):
    u = User.objects.create_user(phone=phone, name=name)
    profile = services.get_or_create_profile(u, origin=engine.ORIGIN_LEGACY, legacy=legacy)
    for k, v in over.items():
        setattr(profile, k, v)
    if over:
        profile.save()
    return u


def лист(response):
    """Первый лист книги из ответа."""
    return load_workbook(BytesIO(response.content)).active


def строки(response):
    """Строки листа без заголовка — кортежами значений."""
    return [tuple(c.value for c in row) for row in лист(response).iter_rows(min_row=2)]


def test_файл_повторяет_колонки_и_порядок_листа(api, params):
    игрок("+7720000001", "Ярмуханбетов Асхат", 12, birth_year=2005, region="Астана")
    игрок("+7720000002", "Абдиров Ерлан Серикович", 55, birth_year=1999, region="Шымкент")
    игрок("+7720000003", "Иванов Пётр", 33.5, birth_year=2012, region="Алматы")

    r = api.get(URL)
    assert r.status_code == 200
    assert r["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    # Дата в имени — по времени сервера (Астана), а не машины, где идёт тест.
    assert r["Content-Disposition"] == 'attachment; filename="rating-%s.xlsx"' % timezone.localdate().isoformat()

    ws = лист(r)
    assert [c.value for c in ws[1]] == ["№", "Рейтинг", "Фамилия Имя Отчество", "Год рождения", "Регион"]
    # По убыванию рейтинга, как лист без сортировки; фамилия прописными.
    assert строки(r) == [
        (1, 55, "АБДИРОВ Ерлан Серикович", 1999, "Шымкент"),
        (2, 33.5, "ИВАНОВ Пётр", 2012, "Алматы"),
        (3, 12, "ЯРМУХАНБЕТОВ Асхат", 2005, "Астана"),
    ]


def test_рейтинг_в_файле_число_с_двумя_знаками(api, params):
    """Число, а не текст «33,50»: по тексту в таблице не посчитать и не
    отсортировать. Два знака задаёт формат ячейки (п. 6.4 Положения)."""
    игрок("+7720000010", "Иванов Пётр", 33.5)

    ячейка = лист(api.get(URL))["B2"]
    assert ячейка.data_type == "n"
    assert ячейка.value == 33.5
    assert ячейка.number_format == "0.00"


def test_отбор_экрана_доходит_до_файла(api, params):
    игрок("+7720000020", "Он Первый", 20, sex="m", region="Астана")
    игрок("+7720000021", "Она Первая", 30, sex="f", region="Алматы")
    игрок("+7720000022", "Она Вторая", 10, sex="f", region="Астана")
    игрок("+7720000023", "Пауза Долгая", 44, sex="f", status=engine.STATUS_INACTIVE)

    def имена(**query):
        return [row[2] for row in строки(api.get(URL, query))]

    # Без отбора — активные; неактивная в текущий лист не входит (п. 18.2).
    assert имена() == ["ОНА Первая", "ОН Первый", "ОНА Вторая"]
    assert имена(sex="f") == ["ОНА Первая", "ОНА Вторая"]
    assert имена(q="Астана") == ["ОН Первый", "ОНА Вторая"]
    assert имена(status=engine.STATUS_INACTIVE) == ["ПАУЗА Долгая"]
    assert имена(all="1") == ["ПАУЗА Долгая", "ОНА Первая", "ОН Первый", "ОНА Вторая"]


def test_сортировка_экрана_доходит_до_файла(api, params):
    игрок("+7720000030", "Ярмуханбетов Асхат", 12)
    игрок("+7720000031", "Абдиров Ерлан", 55)
    игрок("+7720000032", "Иванов Пётр", 33)

    по_имени = строки(api.get(URL, {"sort": "-name"}))
    assert [row[2] for row in по_имени] == ["ЯРМУХАНБЕТОВ Асхат", "ИВАНОВ Пётр", "АБДИРОВ Ерлан"]
    # «№» — порядок в выгруженном списке, а не место в рейтинге.
    assert [row[0] for row in по_имени] == [1, 2, 3]
    assert [row[1] for row in строки(api.get(URL, {"sort": "value"}))] == [12, 33, 55]


def test_в_файл_идёт_весь_отбор_а_не_страница(api, params):
    """Экран показывает сотню строк, файл — все: ради этого его и собирает
    сервер. Параметры страницы на выгрузку не действуют."""
    for i in range(103):
        игрок("+77200001%03d" % i, "Игрок%03d Имя" % i, 10 + i)

    все = строки(api.get(URL, {"page": 2, "page_size": 2}))
    assert len(все) == 103
    assert все[0][:3] == (1, 112, "ИГРОК102 Имя")
    assert все[-1][:3] == (103, 10, "ИГРОК000 Имя")


def test_пустые_год_и_регион_остаются_пустыми_ячейками(api, params):
    """На экране вместо них «—», в таблице — пусто: прочерк в числовой колонке
    ломает сортировку и фильтр по ней."""
    игрок("+7720000040", "Безродный Иван", 20)

    assert строки(api.get(URL)) == [(1, 20, "БЕЗРОДНЫЙ Иван", None, None)]


def test_пустой_отбор_даёт_файл_с_одними_заголовками(api, params):
    игрок("+7720000050", "Кто Угодно", 20, sex="m")

    r = api.get(URL, {"sex": "f"})
    assert r.status_code == 200
    assert лист(r).max_row == 1
    assert лист(r)["C1"].value == "Фамилия Имя Отчество"


def test_выгрузка_открыта_без_входа(api, params):
    """Лист публичный (ТЗ §3, Э0.4) — значит, и его выгрузка: в файле те же
    пять колонок, что видит гость."""
    игрок("+7720000060", "Гость Смотрит", 20)

    assert строки(api.get(URL)) == [(1, 20, "ГОСТЬ Смотрит", None, None)]


@pytest.mark.parametrize(
    "имя, ожидание",
    [
        ("Ахметов Ерлан Серикович", "АХМЕТОВ Ерлан Серикович"),
        ("Әбдіқадыр Нұрлан", "ӘБДІҚАДЫР Нұрлан"),
        ("Ким", "КИМ"),
        ("  Ли Александр ", "ЛИ Александр"),
        # Служебная метка начинается не с буквы и остаётся как есть.
        ("[e2e] Проверочный", "[e2e] Проверочный"),
        ("", ""),
    ],
)
def test_фамилия_прописными_как_на_экране(имя, ожидание):
    assert fio(имя) == ожидание
