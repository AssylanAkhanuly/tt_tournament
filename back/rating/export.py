"""Выгрузка рейтинг-листа в Excel ✳ (01.10.2026).

Здесь только сборка файла: какие строки в него идут и в каком порядке, решает
ручка (`views.RatingExportView`) теми же отбором и сортировкой, что у листа, —
второго правила отбора нет.

Колонки те же пять, что на экране (решение федерации, 15.09.2026): «№»,
«Рейтинг», «Фамилия Имя Отчество», «Год рождения», «Регион».

Рейтинг кладётся числом, а два знака задаёт формат ячейки (п. 6.4 Положения):
текст «33,50» в таблице не посчитать и не отсортировать. Пустые год и регион —
пустые ячейки, а не прочерк, по той же причине.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: Заголовок и ширина колонки (в знаках). Порядок — как в таблице на экране.
COLUMNS = [
    ("№", 7),
    ("Рейтинг", 11),
    ("Фамилия Имя Отчество", 42),
    ("Год рождения", 15),
    ("Регион", 32),
]
VALUE_FORMAT = "0.00"


def fio(name: str) -> str:
    """«Ахметов Ерлан Серикович» → «АХМЕТОВ Ерлан Серикович».

    Фамилия прописными, как на экране (`fio()` в `front/…/rating/lib/format.ts`,
    замечания федерации 15.09.2026) — правило одно, записано в двух местах, и
    менять его надо в обоих. Фамилия — первое слово; строка не с буквы
    (служебная метка) остаётся как есть.
    """
    s = (name or "").strip()
    if not s or not s[0].isalpha():
        return s
    surname = s.split(None, 1)[0]
    return surname.upper() + s[len(surname):]


def rating_xlsx(rows) -> bytes:
    """Книга с рейтинг-листом.

    `rows` — `(значение, ФИО, год рождения, регион)` в том порядке, в каком
    строки идут в листе. «№» — порядок в выгруженном списке, как на экране.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Рейтинг"

    ws.append([title for title, _ in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        ws.cell(row=1, column=i).font = Font(bold=True)
    ws["B1"].alignment = Alignment(horizontal="right")

    for num, (value, name, birth_year, region) in enumerate(rows, start=1):
        ws.append([num, value, fio(name), birth_year, region or None])
        ws.cell(row=num + 1, column=2).number_format = VALUE_FORMAT

    # Заголовок остаётся на месте при прокрутке, у колонок — штатный фильтр.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLUMNS)), ws.max_row)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
